import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from data_processor import (
    COL_ASIGNATURA,
    COL_MATRICULADOS,
    COL_CO_EV_CONT,
    COL_PCT_CO_EV_CONT,
    COL_CO_EV_FIN,
    COL_PCT_CO_EV_FIN,
    COL_CO_TOTAL,
    COL_PCT_CO_TOTAL,
    COL_CE_EV_CONT,
    COL_PCT_CE_EV_CONT,
    COL_CE_EV_FIN,
    COL_PCT_CE_EV_FIN,
    COL_CE_TOTAL,
    COL_PCT_CE_TOTAL,
    COL_TOTAL_ACUMULADO,
    COL_PCT_TOTAL_ACUMULADO,
    get_global_metrics,
)

def generate_styled_excel(df_enriched, degree_name="Titulación Universitaria"):
    """
    Genera un informe Excel maquetado profesionalmente utilizando openpyxl.
    Incluye el nombre de la titulación en el título, subtítulo y nombre de pestaña.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Nombre de la pestaña (límite de 31 caracteres en Excel)
    clean_degree_title = str(degree_name).strip()
    sheet_title = clean_degree_title[:31] if clean_degree_title else "Informe de Notas"
    # Eliminar caracteres no válidos para el nombre de pestaña Excel
    for char in [':', '\\', '/', '?', '*', '[', ']']:
        sheet_title = sheet_title.replace(char, '')
    ws.title = sheet_title if sheet_title else "Informe de Notas"
    
    # Habilitar líneas de cuadrícula
    ws.views.sheetView[0].showGridLines = True
    
    # Estilos
    font_title = Font(name="Calibri", size=15, bold=True, color="1F4E78")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
    
    fill_header_base = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Azul Marino (CO)
    fill_header_ce = PatternFill(start_color="333F48", end_color="333F48", fill_type="solid")   # Gris Oscuro (CE)
    fill_header_tot = PatternFill(start_color="1E4D2B", end_color="1E4D2B", fill_type="solid")  # Verde Bosque (Totales)
    fill_header_matr = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid") # Azul claro
    
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    
    fill_row_even = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_row_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    fill_summary = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_summary = Font(name="Calibri", size=10, bold=True, color="1E4D2B")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_top_thick = Border(top=Side(border_style="medium", color="1E4D2B"), bottom=Side(border_style="double", color="1E4D2B"))

    # Título del Informe con Nombre de Titulación
    ws.merge_cells("A1:P1")
    ws["A1"] = f"INFORME ESTADÍSTICO DE RESULTADOS - {clean_degree_title.upper()}"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    
    ws.merge_cells("A2:P2")
    ws["A2"] = f"Titulación: {clean_degree_title} | Desglose por asignatura, convocatoria (Ordinaria / Extraordinaria) y tipo de evaluación"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 20

    # Columnas
    columns_config = [
        (COL_ASIGNATURA, "Asignatura", fill_header_matr, "@", Alignment(horizontal="left")),
        (COL_MATRICULADOS, "Nº Matriculados", fill_header_matr, "#,##0", Alignment(horizontal="center")),
        
        # Convocatoria Ordinaria
        (COL_CO_EV_CONT, "Aprobados Ev. Cont. (CO)", fill_header_base, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CO_EV_CONT, "% Ev. Cont. (CO)", fill_header_base, "0.0%", Alignment(horizontal="center")),
        (COL_CO_EV_FIN, "Aprobados Ev. Final (CO)", fill_header_base, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CO_EV_FIN, "% Ev. Final (CO)", fill_header_base, "0.0%", Alignment(horizontal="center")),
        (COL_CO_TOTAL, "Total Aprobados (CO)", fill_header_base, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CO_TOTAL, "% Total Aprobados (CO)", fill_header_base, "0.0%", Alignment(horizontal="center")),
        
        # Convocatoria Extraordinaria
        (COL_CE_EV_CONT, "Aprobados Ev. Cont. (CE)", fill_header_ce, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CE_EV_CONT, "% Ev. Cont. (CE)", fill_header_ce, "0.0%", Alignment(horizontal="center")),
        (COL_CE_EV_FIN, "Aprobados Ev. Final (CE)", fill_header_ce, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CE_EV_FIN, "% Ev. Final (CE)", fill_header_ce, "0.0%", Alignment(horizontal="center")),
        (COL_CE_TOTAL, "Total Aprobados (CE)", fill_header_ce, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_CE_TOTAL, "% Total Aprobados (CE)", fill_header_ce, "0.0%", Alignment(horizontal="center")),
        
        # Total Global Acumulado
        (COL_TOTAL_ACUMULADO, "Total Aprobados (CO + CE)", fill_header_tot, "#,##0", Alignment(horizontal="center")),
        (COL_PCT_TOTAL_ACUMULADO, "% Total Aprobados Global", fill_header_tot, "0.0%", Alignment(horizontal="center")),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 28

    for col_idx, (key, label, fill_style, num_fmt, align) in enumerate(columns_config, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=label)
        cell.font = font_header
        cell.fill = fill_style
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell

    start_data_row = 5
    for row_idx, (_, row) in enumerate(df_enriched.iterrows(), start=start_data_row):
        ws.row_dimensions[row_idx].height = 20
        row_fill = fill_row_even if row_idx % 2 == 0 else fill_row_odd
        
        for col_idx, (key, label, fill_style, num_fmt, align) in enumerate(columns_config, start=1):
            val = row[key]
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.fill = row_fill
            cell.number_format = num_fmt
            cell.alignment = align
            cell.border = border_cell

    end_data_row = start_data_row + len(df_enriched) - 1
    summary_row = end_data_row + 1
    ws.row_dimensions[summary_row].height = 24

    metrics = get_global_metrics(df_enriched)

    summary_values = {
        COL_ASIGNATURA: f"TOTALES GLOBALES ({clean_degree_title.upper()})",
        COL_MATRICULADOS: metrics["tot_matriculados"],
        COL_CO_EV_CONT: metrics["tot_co_ev_cont"],
        COL_PCT_CO_EV_CONT: metrics["pct_co_ev_cont"],
        COL_CO_EV_FIN: metrics["tot_co_ev_fin"],
        COL_PCT_CO_EV_FIN: metrics["pct_co_ev_fin"],
        COL_CO_TOTAL: metrics["tot_co_total"],
        COL_PCT_CO_TOTAL: metrics["pct_co_total"],
        COL_CE_EV_CONT: metrics["tot_ce_ev_cont"],
        COL_PCT_CE_EV_CONT: metrics["pct_ce_ev_cont"],
        COL_CE_EV_FIN: metrics["tot_ce_ev_fin"],
        COL_PCT_CE_EV_FIN: metrics["pct_ce_ev_fin"],
        COL_CE_TOTAL: metrics["tot_ce_total"],
        COL_PCT_CE_TOTAL: metrics["pct_ce_total"],
        COL_TOTAL_ACUMULADO: metrics["tot_acumulado"],
        COL_PCT_TOTAL_ACUMULADO: metrics["pct_acumulado"],
    }

    for col_idx, (key, label, fill_style, num_fmt, align) in enumerate(columns_config, start=1):
        val = summary_values[key]
        cell = ws.cell(row=summary_row, column=col_idx, value=val)
        cell.font = font_summary
        cell.fill = fill_summary
        cell.number_format = num_fmt
        cell.border = border_top_thick
        cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "center", vertical="center")

    for col_idx, (key, label, fill_style, num_fmt, align) in enumerate(columns_config, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(label)
        if col_idx == 1:
            for r in range(start_data_row, end_data_row + 1):
                cell_val = str(ws.cell(row=r, column=col_idx).value or "")
                if len(cell_val) > max_len:
                    max_len = len(cell_val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 25), 65)
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
